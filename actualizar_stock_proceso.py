"""
Actualiza 'stock doina.xlsx' con una nueva hoja 'STOCK EN PROCESO'
leyendo la foto más reciente del sistema de reportes web.
"""

import json
import sys
from pathlib import Path
from datetime import datetime, date

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

# ── Rutas ──────────────────────────────────────────────────────────────────────
BASE_DIR    = Path(__file__).resolve().parent
EXCEL_PATH  = BASE_DIR / "stock doina.xlsx"
FOTOS_DIR   = BASE_DIR.parent / "proyecto reportes web doina" / "data" / "fotos"
NOMBRE_HOJA = "STOCK EN PROCESO"

# ── Colores ────────────────────────────────────────────────────────────────────
COLOR_TITULO   = "111318"   # negro marca
COLOR_HEADER   = "1a1c23"   # gris oscuro
COLOR_NARANJA  = "e87722"   # naranja DOINA
COLOR_VENCIDO  = "FFF5F5"   # rosa claro para vencidos
COLOR_ALERTA   = "FFFBE6"   # amarillo claro para alertas
COLOR_HEADER_F = "FFFFFF"   # texto blanco en header
COLOR_TITLE_F  = "FFFFFF"   # texto blanco en título
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def cargar_ultima_foto() -> dict:
    archivos = sorted(FOTOS_DIR.glob("*.json"), reverse=True)
    if not archivos:
        sys.exit("ERROR: No hay fotos en data/fotos/")
    ruta = archivos[0]
    print(f"Leyendo foto: {ruta.name}")
    with open(ruta, encoding="utf-8") as f:
        return json.load(f)


def extraer_lotes_activos(foto: dict) -> list[dict]:
    """Devuelve todos los lotes activos de todas las etapas."""
    resultado = []
    for etapa_key, datos in foto.get("etapas", {}).items():
        nombre_etapa = datos.get("nombre", etapa_key)
        for lote in datos.get("lotes", []):
            if lote.get("estado") != "activo":
                continue
            resultado.append({
                "etapa":        nombre_etapa,
                "lote":         lote.get("lote", ""),
                "grupo":        lote.get("grupo_mercaderia", ""),
                "descripcion":  lote.get("descripcion_mercaderia", ""),
                "formato":      lote.get("formato") or "",
                "fecha_entrada": lote.get("fecha_entrada"),
                "dias":         lote.get("dias_parciales"),
                "unidades":     lote.get("unidades"),
                "kg_netos":     lote.get("kg_netos"),
                "vencimiento":  lote.get("vencimiento"),
                "vencido":      lote.get("vencido", False),
                "alerta_dias":  lote.get("alerta_dias", False),
                "alerta_merma": lote.get("alerta_merma", False),
                "merma_pct":    lote.get("merma_total_pct"),
            })
    # Vencidos primero, luego por días desc
    resultado.sort(key=lambda l: (not l["vencido"], -(l["dias"] or 0)))
    return resultado


def set_col_width(ws, col: int, width: float):
    ws.column_dimensions[get_column_letter(col)].width = width


def crear_hoja_stock_proceso(wb: openpyxl.Workbook, foto: dict):
    # Eliminar hoja si ya existe
    if NOMBRE_HOJA in wb.sheetnames:
        del wb[NOMBRE_HOJA]

    ws = wb.create_sheet(NOMBRE_HOJA, 0)   # primera posición

    fecha_foto = foto.get("fecha_foto", "")
    resumen    = foto.get("resumen", {})
    lotes      = extraer_lotes_activos(foto)
    total_lotes = len(lotes)
    total_kg    = sum(l["kg_netos"] or 0 for l in lotes)
    vencidos    = sum(1 for l in lotes if l["vencido"])

    # ── Fila 1: Título ─────────────────────────────────────────────────────────
    ws.merge_cells("A1:N1")
    c = ws["A1"]
    c.value     = f"STOCK EN PROCESO — DOINA  |  Foto: {fecha_foto}"
    c.font      = Font(name="Calibri", bold=True, size=14, color=COLOR_TITLE_F)
    c.fill      = PatternFill("solid", fgColor=COLOR_TITULO)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    # ── Fila 2: KPIs rápidos ───────────────────────────────────────────────────
    ws.row_dimensions[2].height = 22
    kpis = [
        ("A2:C2", f"Lotes activos: {total_lotes:,}"),
        ("D2:F2", f"Kg en proceso: {total_kg:,.0f}"),
        ("G2:I2", f"Vencidos: {vencidos}"),
        ("J2:N2", f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"),
    ]
    for rng, txt in kpis:
        ws.merge_cells(rng)
        start = rng.split(":")[0]
        c = ws[start]
        c.value     = txt
        c.font      = Font(name="Calibri", size=10, color="AAAAAA", italic=True)
        c.fill      = PatternFill("solid", fgColor="16181E")
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # ── Fila 3: vacía ──────────────────────────────────────────────────────────
    ws.row_dimensions[3].height = 6

    # ── Fila 4: Encabezados ────────────────────────────────────────────────────
    HEADERS = [
        ("ETAPA",         18),
        ("LOTE",          13),
        ("GRUPO",         22),
        ("DESCRIPCION",   28),
        ("FORMATO",       12),
        ("F. ENTRADA",    13),
        ("DIAS",          7),
        ("UNIDADES",      10),
        ("KG NETOS",      11),
        ("VENCIMIENTO",   13),
        ("VENCIDO",       9),
        ("ALERTA DIAS",   11),
        ("ALERTA MERMA",  12),
        ("MERMA %",       10),
    ]
    ws.row_dimensions[4].height = 22
    for col, (header, width) in enumerate(HEADERS, start=1):
        c = ws.cell(row=4, column=col, value=header)
        c.font      = Font(name="Calibri", bold=True, size=9, color=COLOR_HEADER_F)
        c.fill      = PatternFill("solid", fgColor=COLOR_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = BORDER
        set_col_width(ws, col, width)

    # ── Filas de datos ─────────────────────────────────────────────────────────
    DATE_FMT = "DD/MM/YYYY"
    for row_idx, l in enumerate(lotes, start=5):
        ws.row_dimensions[row_idx].height = 16

        # Color de fondo de fila
        if l["vencido"]:
            bg = PatternFill("solid", fgColor="FFF0F0")
        elif l["alerta_dias"] or l["alerta_merma"]:
            bg = PatternFill("solid", fgColor="FFFBE6")
        else:
            bg = PatternFill("solid", fgColor="FFFFFF") if row_idx % 2 == 0 else PatternFill("solid", fgColor="F8F9FB")

        valores = [
            l["etapa"],
            l["lote"],
            l["grupo"],
            l["descripcion"],
            l["formato"],
            l["fecha_entrada"],    # fecha (str YYYY-MM-DD)
            l["dias"],
            l["unidades"],
            l["kg_netos"],
            l["vencimiento"],      # fecha (str YYYY-MM-DD)
            "SI" if l["vencido"] else "",
            "SI" if l["alerta_dias"] else "",
            "SI" if l["alerta_merma"] else "",
            l["merma_pct"],
        ]
        alineaciones = [
            "left", "left", "left", "left", "center",
            "center", "right", "right", "right",
            "center", "center", "center", "center", "right",
        ]
        for col, (val, alin) in enumerate(zip(valores, alineaciones), start=1):
            c = ws.cell(row=row_idx, column=col)
            c.border    = BORDER
            c.fill      = bg
            c.alignment = Alignment(horizontal=alin, vertical="center")
            c.font      = Font(name="Calibri", size=9)

            # Convertir fechas str → datetime para que Excel las muestre bien
            if col in (6, 10) and isinstance(val, str) and val:
                try:
                    val = datetime.strptime(val, "%Y-%m-%d")
                    c.number_format = DATE_FMT
                except ValueError:
                    pass

            c.value = val

            # Resaltar vencido
            if col == 11 and l["vencido"]:
                c.font = Font(name="Calibri", size=9, bold=True, color="CC0000")

            # Resaltar merma %
            if col == 14 and val is not None:
                c.number_format = "0.00%"
                # Convertir a fracción si viene como porcentaje >1
                if isinstance(val, (int, float)) and val > 1:
                    c.value = val / 100

    # ── Fila totales ───────────────────────────────────────────────────────────
    total_row = 5 + len(lotes)
    ws.row_dimensions[total_row].height = 18
    ws.merge_cells(f"A{total_row}:G{total_row}")
    c = ws[f"A{total_row}"]
    c.value     = f"TOTAL — {total_lotes} lotes activos"
    c.font      = Font(name="Calibri", bold=True, size=10, color=COLOR_TITLE_F)
    c.fill      = PatternFill("solid", fgColor=COLOR_NARANJA)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.border    = BORDER

    for col, val in [(8, sum(l["unidades"] or 0 for l in lotes)),
                     (9, total_kg)]:
        c = ws.cell(row=total_row, column=col)
        c.value     = val
        c.font      = Font(name="Calibri", bold=True, size=10, color=COLOR_TITLE_F)
        c.fill      = PatternFill("solid", fgColor=COLOR_NARANJA)
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border    = BORDER
        c.number_format = "#,##0"

    # Celdas vacías del total row
    for col in list(range(10, 15)):
        c = ws.cell(row=total_row, column=col)
        c.fill   = PatternFill("solid", fgColor=COLOR_NARANJA)
        c.border = BORDER

    # ── Congelar encabezado ────────────────────────────────────────────────────
    ws.freeze_panes = "A5"

    # ── Auto-filter ───────────────────────────────────────────────────────────
    ws.auto_filter.ref = f"A4:N{4 + len(lotes)}"

    print(f"  Hoja '{NOMBRE_HOJA}' creada con {len(lotes)} lotes activos ({total_kg:,.0f} kg)")
    return ws


def main():
    if not EXCEL_PATH.exists():
        sys.exit(f"ERROR: No se encontró {EXCEL_PATH}")

    foto = cargar_ultima_foto()
    wb   = openpyxl.load_workbook(str(EXCEL_PATH))

    crear_hoja_stock_proceso(wb, foto)

    wb.save(str(EXCEL_PATH))
    print(f"Excel guardado: {EXCEL_PATH.name}")


if __name__ == "__main__":
    main()
